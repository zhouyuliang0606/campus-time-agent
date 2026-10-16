"""第二批测试：文件上传、边界防护、API 配置（含脱敏）、通知附件、Office 文档。

人话：这一组测的是"喂给 AI 的东西"——学生能不能上传课表、
乱传文件会不会被拦、管理员填的密钥会不会被泄露回去、
Word/Excel/PPT 里的内容（特别是表格）能不能真被读出来。
"""
import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _harness import Checker, data_path, make_client, sandbox, title  # noqa: E402

# 这一组会保存假密钥、上传临时文件，全部关进沙箱里跑
with sandbox():
    client = make_client()
    c = Checker()

    def up(name, content, ctype="application/octet-stream"):
        """模拟在网页上选一个文件上传。"""
        return client.post("/api/upload", files={"file": (name, content, ctype)})

    title("1. 文件上传")
    r = up("课程表.txt", "周一 高等数学 8:00\n周二 线性代数 10:00".encode("utf-8"))
    f1 = r.json().get("file", {})
    c.check("上传 UTF-8 文本", r.status_code == 200 and f1.get("id"), f"id={f1.get('id')} chars={f1.get('chars')}")
    c.check("元信息含原文件名", f1.get("name") == "课程表.txt")
    c.check("保留了原件(可下载)", f1.get("has_raw") is True)

    r = up("gbk.txt", "这是一段GBK编码的中文".encode("gbk"))
    f2 = r.json().get("file", {})
    c.check("上传 GBK 中文文本", r.status_code == 200 and f2.get("chars") > 0, f"chars={f2.get('chars')}")

    title("2. 上传的边界防护")
    c.check("空文件被拒 400", up("empty.txt", b"").status_code == 400)
    c.check("不支持的类型被拒 415", up("virus.exe", b"MZ\x00\x00").status_code == 415)
    c.check("超限文件被拒 413", up("big.txt", b"x" * (5 * 1024 * 1024 + 1)).status_code == 413)

    title("3. 文件列表 / 读取 / 下载")
    lst = client.get("/api/uploads").json()["files"]
    c.check("列表有记录", len(lst) >= 2, f"共 {len(lst)} 个")
    c.check("列表倒序(新的在前)", lst[0]["name"] == "gbk.txt", f"首位={lst[0]['name']}")
    c.check("读取提取的文字", "高等数学" in client.get(f"/api/uploads/{f1['id']}").json()["content"])
    c.check("读不存在的文件 404", client.get("/api/uploads/nope123").status_code == 404)
    d = client.get(f"/api/uploads/{f1['id']}/download")
    c.check("下载原件成功", d.status_code == 200 and "8:00" in d.text)

    title("3b. 上传的文件必须落在**当前数据目录**里（沙箱隔离，别污染演示环境）")
    # 2026-10-16 实测到的真事：原来 `UPLOAD_DIR = os.path.join(DATA_DIR, "uploads")` 在导入时就定死，
    # 沙箱把 CAMPUSTIME_DATA_DIR 换了目录，上传却仍然写进真实 `app/data/uploads/` ——
    # 两轮全量往里堆了 30 个测试文件才被发现。`_harness` 的 SCRATCH 一直列着 "uploads"，
    # 也就是说沙箱从头就假设这件事是隔离好的，只有写入方没做到（与 #47 那条账本没隔离同根）。
    import app.store as S
    inside = os.path.exists(data_path("uploads", f1["id"] + ".txt"))
    outside = os.path.exists(os.path.join(S._DEFAULT_DATA_DIR, "uploads", f1["id"] + ".txt"))
    c.check("正文落在沙箱里", inside, data_path("uploads", f1["id"] + ".txt"))
    c.check("真实 app/data/uploads/ 里**没有**这个测试文件（有就说明上传目录又漂回导入时那个值了）",
            not outside, os.path.join(S._DEFAULT_DATA_DIR, "uploads", f1["id"] + ".txt"))
    # 带原件的那一条（gbk.txt 有原件可下载）：原件也不许漏到真实目录
    real_up_dir = os.path.join(S._DEFAULT_DATA_DIR, "uploads")
    stray_raw = [f for f in (os.listdir(real_up_dir) if os.path.isdir(real_up_dir) else [])
                 if f.startswith(f2["id"])]
    c.check("带原件的上传：原件落在沙箱里、真实目录里一个都没有",
            not stray_raw and os.path.exists(data_path("uploads", f2["id"] + ".txt")),
            f"真实目录漏出：{stray_raw}")
    src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "app", "store.py"),
               encoding="utf-8").read()
    src_code = "\n".join(ln.split("#")[0] for ln in src.splitlines())
    c.check("形状：拼上传路径只许走 _upload_dir()（每次现读 env），不许再用导入时那个 UPLOAD_DIR",
            "os.path.join(_upload_dir()" in src_code and "os.path.join(UPLOAD_DIR" not in src_code,
            f"_upload_dir {src_code.count('os.path.join(_upload_dir()')} 处 / "
            f"UPLOAD_DIR {src_code.count('os.path.join(UPLOAD_DIR')} 处")

    title("4. API 配置读取与脱敏")
    s = client.get("/api/admin/settings").json()
    c.check("读取配置", "base_url" in s["settings"])
    c.check("未配置时 configured=False", s["configured"] is False)

    FAKE = "sk-abcdefghijklmnop1234"
    r = client.post("/api/admin/settings", json={"api_key": FAKE, "model": "deepseek-test"})
    c.check("保存密钥+模型", r.status_code == 200 and r.json()["ok"])
    s2 = client.get("/api/admin/settings").json()
    c.check("配置来源已变为管理后台", s2["source"] == "管理后台配置", s2["source"])
    masked = s2["settings"]["api_key"]
    c.check("已配置 configured=True", s2["configured"] is True)
    c.check("返回的密钥被打码", "*" in masked and FAKE not in masked, f"掩码={masked}")
    c.check("模型已生效", s2["settings"]["model"] == "deepseek-test")

    title("5. 防呆：不要把掩码串当成真密钥存回去")
    client.post("/api/admin/settings", json={"api_key": masked, "model": "keep-me"})
    saved = json.load(open(data_path("settings.json"), encoding="utf-8"))
    c.check("掩码串没有被写入", saved.get("api_key") == FAKE, f"落盘值={str(saved.get('api_key'))[:10]}...")
    c.check("其他字段仍正常保存", saved.get("model") == "keep-me")

    client.post("/api/admin/settings", json={"model": ""})  # 空字符串 = "这一项不配置了"
    c.check("空字符串会回落到默认模型",
            client.get("/api/admin/settings").json()["settings"]["model"] == "deepseek-chat")

    # 把配置清空，模拟"管理员还没填过"的状态：
    # 这里写成空对象而不是删文件，因为 store 每次现读，写空就等于没配过
    with open(data_path("settings.json"), "w", encoding="utf-8") as f:
        f.write("{}")

    title("6. 连通性自检（无密钥时快速失败，不发网络请求）")
    r = client.post("/api/admin/settings/test")
    c.check("无密钥时提示先配置", r.status_code == 200 and r.json()["ok"] is False, r.json().get("message", "")[:40])

    title("7. 通知支持附件")
    r = up("放假安排.txt", "国庆放假 10/1-10/7".encode("utf-8"))
    att = {"id": r.json()["file"]["id"], "name": r.json()["file"]["name"]}
    r = client.post("/api/admin/notify", json={"title": "国庆放假通知", "content": "详见附件", "attachment": att})
    c.check("发布带附件的通知", r.status_code == 200 and r.json()["notice"].get("attachment"))
    r = client.post("/api/admin/notify", json={"title": "", "content": "", "attachment": {"id": "ghost123"}})
    c.check("幽灵附件被拒 400", r.status_code == 400)
    c.check("标题内容全空被拒 400", client.post("/api/admin/notify", json={"title": "", "content": ""}).status_code == 400)
    ns = client.get("/api/notices").json()["notices"]
    c.check("通知列表带附件信息", any(n.get("attachment") for n in ns))

    title("8. 新版 Office 文档上传（Word/Excel/PPT，含表格）")
    try:
        import docx
        from openpyxl import Workbook
        from pptx import Presentation

        # 带标题 + 正文 + 表格的 Word（课表正是这种结构）
        d = docx.Document()
        d.add_heading("2026 秋季课表", level=1)
        d.add_paragraph("软件工程 25 级 3 班")
        tb = d.add_table(rows=2, cols=3)
        for (r_, c_), v in {(0, 0): "星期", (0, 1): "课程", (0, 2): "时间",
                            (1, 0): "周一", (1, 1): "高等数学", (1, 2): "8:00"}.items():
            tb.cell(r_, c_).text = v
        b1 = io.BytesIO(); d.save(b1); docx_bytes = b1.getvalue()

        # 两个工作表的 Excel
        wb = Workbook(); ws = wb.active; ws.title = "待取清单"
        ws.append(["取件码", "收件人", "滞留天数"]); ws.append(["B-5678", "李四", 5])
        w2 = wb.create_sheet("统计"); w2.append(["项目", "数量"]); w2.append(["滞留", 3])
        b2 = io.BytesIO(); wb.save(b2); xlsx_bytes = b2.getvalue()

        # 一页 PPT，标题 + 正文
        prs = Presentation()
        s1 = prs.slides.add_slide(prs.slide_layouts[1])
        s1.shapes.title.text = "快递高峰应对方案"
        s1.placeholders[1].text = "延长营业时间"
        b3 = io.BytesIO(); prs.save(b3); pptx_bytes = b3.getvalue()

        r = up("我的课表.docx", docx_bytes)
        ok = r.status_code == 200
        c.check("上传 Word .docx", ok, f"提取 {r.json().get('file', {}).get('chars')} 字")
        if ok:
            body = client.get(f"/api/uploads/{r.json()['file']['id']}").json()["content"]
            c.check("Word 的表格内容被读到", "高等数学" in body and "星期" in body, f"片段={body[:36]!r}")

        r = up("驿站清单.xlsx", xlsx_bytes)
        ok = r.status_code == 200
        c.check("上传 Excel .xlsx", ok, f"提取 {r.json().get('file', {}).get('chars')} 字")
        if ok:
            body = client.get(f"/api/uploads/{r.json()['file']['id']}").json()["content"]
            c.check("Excel 两个工作表都读到", "待取清单" in body and "统计" in body and "B-5678" in body)

        r = up("方案.pptx", pptx_bytes)
        ok = r.status_code == 200
        c.check("上传 PPT .pptx", ok, f"提取 {r.json().get('file', {}).get('chars')} 字")
        if ok:
            body = client.get(f"/api/uploads/{r.json()['file']['id']}").json()["content"]
            c.check("PPT 标题正文都读到", "快递高峰应对方案" in body and "延长营业时间" in body)

        title("9. 老版二进制格式要给「另存为」的可执行指引")
        r = up("旧课表.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
        msg = r.json().get("error", "") if r.status_code != 200 else ""
        c.check("老版 .doc 提示另存为 .docx", r.status_code == 415 and "另存为" in msg and ".docx" in msg, msg[:40])
        r = up("旧表格.xls", b"\xd0\xcf\x11\xe0" + b"\x00" * 64)
        err = r.json().get("error", "")
        c.check("老版 .xls 提示另存为 .xlsx", "另存为" in err and ".xlsx" in err)
    except ImportError as e:
        print(f"  ⚠️ 跳过 Office 相关用例（缺依赖：{e}）")

    sys.exit(c.summary("第二批（上传与配置）"))
