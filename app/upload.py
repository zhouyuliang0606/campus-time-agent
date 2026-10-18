"""上传文件解析（人话：把用户传上来的各种格式，统一变成"一串 AI 能读的纯文字"）。

为什么单独拎成一个文件？
因为"读文件"这件事跟"存文件"（store.py）、"对外接口"（main.py）都不是一回事，
单独放一层，以后想加新格式，只改这里就行，别的地方不用动。

目前支持四类：
  1. 纯文本类（txt/md/csv/json…）→ 按编码直接解码
  2. PDF                        → pypdf 逐页抽
  3. Word / Excel / PPT 新版格式 → python-docx / openpyxl / python-pptx
  4. 老版 Office 二进制格式      → 不支持，给"另存为"的明确指引

变成文字之后还管一件事：把长文档**切成带位置的小块**（chunk_text，见文件下半段）。
不切块的话，几万字的手册只能读开头那一截，后半本等于没上传。
"""
import io
import os

# 明确支持的纯文本后缀（照 md/csv 这类直接按文本读就行）
TEXT_EXT = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".log",
    ".yml", ".yaml", ".xml", ".html", ".htm", ".ini", ".cfg",
    ".py", ".js", ".java", ".c", ".cpp", ".sql", ".rst",
}

# 新版 Office 格式（本质是 OOXML 压缩包，各自需要一个小库来读）
OFFICE_EXT = {".docx", ".xlsx", ".pptx"}

# 老版 Office 二进制格式（.doc/.xls/.ppt），新版库都读不了，需要专门提示用户转换
LEGACY_EXT = {".doc", ".xls", ".ppt"}
LEGACY_NEW = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}

# Excel 最多读多少行：超大表格会把内存和 token 都吃光，宁可截断
MAX_EXCEL_ROWS = 2000


class UploadError(Exception):
    """上传文件处理失败（人话：这文件搞不定，原因写给用户看）。"""


def _ext(name: str) -> str:
    """取文件名后缀，统一转小写方便比较。"""
    return os.path.splitext(name or "")[1].lower()


def ext_supported(name: str) -> bool:
    """这个文件后缀咱们能不能处理。"""
    return _ext(name) in (TEXT_EXT | OFFICE_EXT | {".pdf"})


def supported_desc() -> str:
    """给前端/报错文案用的一句话说明。"""
    return "txt / md / csv / json 等文本文件、PDF，以及 Word(.docx) / Excel(.xlsx) / PPT(.pptx)"


def unsupported_reason(name: str) -> str:
    """说明某个后缀为什么不支持（人话：能指条明路，就别只说"不行"）。

    老版 Office 格式值得单独给一句——用户只要去 Word 里「另存为 .docx」就能解决，
    比笼统一句"暂不支持"有用得多。
    """
    ext = _ext(name)
    if ext in LEGACY_EXT:
        return (
            f"{name} 是老版 Office 二进制格式，读不了。"
            f"请用 Word / WPS / Excel 打开后「另存为」{LEGACY_NEW[ext]}，再重新上传。"
        )
    return f"暂不支持 {name} 这种类型，目前支持：{supported_desc()}"


def extract_text(filename: str, data: bytes) -> str:
    """把一个上传文件变成纯文本。

    :param filename: 原始文件名（靠后缀判断怎么解析）
    :param data: 文件原始字节
    :return: 提取出的文字
    """
    ext = _ext(filename)

    # 老版 Office 二进制格式：新版库都打不开，与其甩底层报错，
    # 不如直接告诉用户"另存为新版后缀就行"（他只要照做就能解决）
    if ext in LEGACY_EXT:
        raise UploadError(
            f"{filename} 是老版 Office 格式，读不了。"
            f"请用 Word/WPS/Excel 打开后「另存为」{LEGACY_NEW[ext]}，再重新上传。"
        )

    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext == ".xlsx":
        text = _xlsx_text(data)
    elif ext == ".pptx":
        text = _pptx_text(data)
    else:
        text = _plain_text(data)

    text = text.strip()
    if not text:
        scan = (
            "如果是扫描版 PDF（整页都是图片），需要先做 OCR 识别，本作品暂不支持。"
            if ext == ".pdf"
            else "如果是图片型文档/扫描件，需要先做 OCR，本作品暂不支持。"
        )
        raise UploadError(f"没能从 {filename} 里读到任何文字。{scan}")
    return text


# ============ 分块（长文档要能被检索，靠的不是一次读完） ============
# 为什么要有这一步：一份几十页的手册有几万字，一次性喂给模型既撑爆上下文也白烧 token；
# 而原先的做法更糟——超过 6000 字直接砍掉后半本，等于后半本**没上传**（#40 修的就是它）。
# 切成带位置的小块之后，回答问题只需要捞出相关的那几块。
#
# 已知限制照实说：PDF / docx 的**表格**抽取本来就弱（见下面几个 _*_text 函数），
# 表格行常常挤成一长串没有标点的话，这种内容切出来会比正文碎，检索命中也差一些。
CHUNK_SIZE = 600           # 目标块大小（字）
CHUNK_OVERLAP = 80         # 相邻块重叠的字数：跨边界的那句话两块都留着，检索才找得着
_MIN_KEEP = CHUNK_SIZE // 2      # 一块至少留这么多字：段落一密就切出两三百字的碎块没意义
_TOLERANCE = CHUNK_SIZE // 2     # 允许超出目标多少去够下一个边界

# 句末与次级停顿（中英文标点都认，文档里两种混着用很常见）
_SENT_END = "。！？!?；;.\n"
_SOFT_END = "，,、：:）)】」』”’ \t"


def _para_end(text: str, i: int) -> bool:
    """i 是不是「空行段落」的末尾：往前数连续空白里至少有两个换行。

    认 `\n \n` 这种中间夹空格的空行——从 PDF / 网页抽出来的文本经常长这样。
    """
    j, lines = i - 1, 0
    while j >= 0 and text[j] in " \t\n":
        if text[j] == "\n":
            lines += 1
        j -= 1
    return lines >= 2


def _cut_at(text: str, start: int) -> int:
    """从 start 起该切到哪儿（返回不含的结束位置）。

    优先级：空行段落 → 句末标点 → 次级停顿 → 硬切。
    三档都是**从后往前**扫，取窗口里最靠后的那个边界：块大小因此贴着 CHUNK_SIZE，
    不会因为文档段落密就碎成几十个字一块（块太碎 → 检索时每块上下文太少，命中也差）。
    """
    n = len(text)
    if start + CHUNK_SIZE >= n:
        return n
    lo = start + _MIN_KEEP
    hi = min(start + CHUNK_SIZE + _TOLERANCE, n)
    for i in range(hi, lo, -1):
        if _para_end(text, i):
            return i
    for i in range(hi, lo, -1):
        if text[i - 1] in _SENT_END:
            return i
    for i in range(hi, lo, -1):
        if text[i - 1] in _SOFT_END:
            return i
    return start + CHUNK_SIZE


def chunk_text(text: str) -> list[dict]:
    """把整篇文字切成带位置的小块（人话：手册拆成一页页卡片，每张记住自己在第几页）。

    返回 `[{"i": 0, "start": 0, "text": "…"}]`，三个字段各有用处：
      - `text` 是原文的**逐字切片**，不改写、不 trim、不压缩空白 —— 拼回去必须还是原文
      - `start` 是它在原文里的字符偏移，将来要说"出自第几块、原文哪一段"靠它
      - `i` 是块号，检索结果里给模型当出处引用
    相邻块重叠 CHUNK_OVERLAP 个字（是故意的，见上面那个常量旁边的注释）。
    """
    text = text or ""
    n = len(text)
    if n == 0:
        return []
    chunks: list[dict] = []
    start = 0
    while True:
        end = _cut_at(text, start)
        chunks.append({"i": len(chunks), "start": start, "text": text[start:end]})
        if end >= n:
            break
        nxt = end - CHUNK_OVERLAP
        # 块短到装不下重叠量时（只有末尾那一块会这样）就直接接上，否则会原地打转
        start = nxt if nxt > start else end
    return chunks


def _plain_text(data: bytes) -> str:
    """纯文本文件：按常见中文/国际编码依次试，都失败就忽略非法字符硬解。"""
    for enc in ("utf-8", "gbk", "gb18030", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    # 都解不出来也不能让用户卡住，至少把能认的字拿出来
    return data.decode("utf-8", errors="ignore")


def _pdf_text(data: bytes) -> str:
    """PDF：用 pypdf 逐页抽取文字。"""
    try:
        from pypdf import PdfReader
    except ImportError:  # 没装依赖时给个明确指引，别甩一堆看不懂的报错
        raise UploadError("读取 PDF 需要额外依赖，请先执行：pip install pypdf") from None

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    """Word(.docx)：把段落和表格都抽出来。

    表格很容易被忽略，但课表、名单这类校园资料恰恰大量用表格，
    所以必须一起读，不能只读 paragraphs。
    """
    try:
        import docx  # python-docx
    except ImportError:
        raise UploadError("读取 Word 需要额外依赖：pip install python-docx") from None

    doc = docx.Document(io.BytesIO(data))
    parts = []
    # 正文段落（含标题）
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    # 表格：一行拼成一段，单元格之间用制表符分隔，AI 仍能看出列的含义
    for i, table in enumerate(doc.tables, 1):
        parts.append(f"【表格 {i}】")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    """Excel(.xlsx)：逐个工作表、逐行转成文本。

    data_only=True 表示取公式**算出来的结果**而不是公式本身——
    用户想看的是数值，不是 =SUM(A1:A10) 这种字符串。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise UploadError("读取 Excel 需要额外依赖：pip install openpyxl") from None

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    truncated = False
    for ws in wb.worksheets:
        parts.append(f"【工作表：{ws.title}】")
        for n, row in enumerate(ws.iter_rows(values_only=True), 1):
            if n > MAX_EXCEL_ROWS:
                truncated = True
                break
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                parts.append("\t".join(cells))
    if truncated:
        parts.append(f"（工作表内容较多，每个表最多展示前 {MAX_EXCEL_ROWS} 行）")
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str:
    """PPT(.pptx)：逐页抽出文字框和表格。"""
    try:
        from pptx import Presentation  # python-pptx
    except ImportError:
        raise UploadError("读取 PPT 需要额外依赖：pip install python-pptx") from None

    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            # 正文/标题文本框
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    lines.append(t)
            # 幻灯片里的表格
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append("\t".join(cells))
        if lines:
            parts.append(f"【第 {i} 页】\n" + "\n".join(lines))
    return "\n".join(parts)
